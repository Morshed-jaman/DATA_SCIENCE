"""Independent Step 7 checks; run with the project .venv from any directory."""
from pathlib import Path
import csv, math, statistics, json, re, hashlib, warnings
from collections import Counter
from zipfile import ZipFile
from lxml import etree
import numpy as np
import pandas as pd
from scipy import stats
import nbformat
import pymupdf
from docx import Document

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'audit'
checks = []
def close(label, actual, expected):
    np.testing.assert_allclose(actual, expected, rtol=1e-10, atol=0, err_msg=label)
    checks.append(label)
def load(name):
    return pd.read_csv(ROOT/'results'/name, float_precision='round_trip')
rows = list(csv.DictReader((ROOT/'ui_ab_testing_data.csv').open(newline='')))
groups = {g:[float(r['session_duration_seconds']) for r in rows if r['ui_version']==g] for g in ['UI_A','UI_B']}
a,b = groups.values()
nA,nB = len(a),len(b)
mA,mB = statistics.mean(a),statistics.mean(b)
sA,sB = statistics.stdev(a),statistics.stdev(b)
desc = load('descriptive_statistics.csv').set_index('ui_version')
normal = load('normality_assessment.csv').set_index('ui_version')
expected_values = []
for g,x in groups.items():
    vals=[len(x),statistics.mean(x),statistics.stdev(x),min(x),*np.quantile(x,[.25,.5,.75]),max(x)]
    close(g+' descriptive statistics',desc.loc[g].to_numpy(float), vals)
    expected_values.extend(vals)
    with warnings.catch_warnings(record=True) as caught:
        warnings.simplefilter('always'); w,p=stats.shapiro(x)
    close(g+' normality',normal.loc[g,['sample_skewness','shapiro_W','shapiro_p_value']].to_numpy(float),[stats.skew(x,bias=False),w,p])
    assert p>0
    assert bool(caught)==(g=='UI_A')
    expected_values += [stats.skew(x,bias=False),w,p]
# Median-centered Levene computed as ANOVA on absolute median deviations.
za=np.abs(np.array(a)-statistics.median(a)); zb=np.abs(np.array(b)-statistics.median(b))
grand=(sum(za)+sum(zb))/(nA+nB)
lev=(nA+nB-2)*(nA*(za.mean()-grand)**2+nB*(zb.mean()-grand)**2)/(sum((za-za.mean())**2)+sum((zb-zb.mean())**2))
lp=stats.f.sf(lev,1,nA+nB-2)
close('manual median Levene',load('variance_test.csv')[['statistic','p_value']].iloc[0],[lev,lp])
vA,vB=sA*sA/nA,sB*sB/nB
delta=mB-mA; se=math.sqrt(vA+vB)
df=(vA+vB)**2/(vA*vA/(nA-1)+vB*vB/(nB-1))
t=delta/se; p=2*stats.t.sf(abs(t),df); margin=stats.t.ppf(.975,df)*se
lo,hi=delta-margin,delta+margin
assert p>0 and lp>0
test=load('t_test_results.csv').iloc[0]
for key,value in dict(n_A=nA,n_B=nB,mean_A=mA,mean_B=mB,sample_sd_A=sA,sample_sd_B=sB,mean_difference_seconds=delta,standard_error_seconds=se,t_statistic=t,degrees_of_freedom=df,p_value=p,ci_95_low_seconds=lo,ci_95_high_seconds=hi,manual_t=t,manual_df=df,manual_p_value=p,manual_ci_low=lo,manual_ci_high=hi).items():
    close(key,test[key],value)
assert test.difference_direction=='UI_B minus UI_A' and test.test_name=='Welch'
expected_values += [lev,lp,delta,se,df,t,p,lo,hi,vA,vB,stats.t.ppf(.975,df)]
probs=load('probability_scenarios.csv')
probability_results=[]
for i,(x,lower,upper) in enumerate([(a,mB,math.inf),(a,-math.inf,80),(b,150,math.inf),(a,80,120)]):
    mean,sd=statistics.mean(x),statistics.stdev(x)
    zl,zu=(lower-mean)/sd,(upper-mean)/sd
    # Independent CDF via erfc instead of scipy.stats.norm.
    cdf=lambda z: .5*math.erfc(-z/math.sqrt(2))
    prob=.5*math.erfc(zl/math.sqrt(2)) if upper==math.inf else cdf(zu)-cdf(zl)
    count=sum(lower<v<upper for v in x); n=len(x); observed=count/n
    for key,value in dict(mean_seconds=mean,sample_sd_seconds=sd,lower_bound_seconds=lower,upper_bound_seconds=upper,z_lower=zl,z_upper=zu,normal_probability=prob,normal_percentage=100*prob,empirical_count=count,denominator=n,empirical_proportion=observed,empirical_percentage=100*observed,empirical_minus_normal_percentage_points=100*(observed-prob)).items():
        close(f'probability {i+1} {key}',probs.iloc[i][key],value)
        if math.isfinite(value): expected_values.append(value)
    probability_results.append({'event':i+1,'probability':prob,'count':count,'denominator':n})
increase=100*delta/mA
business=load('business_impact.csv').iloc[0]
for key,value in dict(mean_A_seconds=mA,mean_B_seconds=mB,absolute_difference_seconds=delta,percentage_increase=increase,p_value=p,ci_95_low_seconds=lo,ci_95_high_seconds=hi).items(): close('business '+key,business[key],value)
expected_values.append(increase)

def xml_text(path):
    with ZipFile(path) as z: root=etree.fromstring(z.read('word/document.xml'))
    ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    return '\n'.join(''.join(e.itertext()) for e in root.xpath('//w:t|//m:t',namespaces=ns))
original=Document(ROOT/'A_B_Testing_Hypothesis_Testing_Assignment.docx')
assert len(original.tables)==2
(OUT/'original_assignment_actual.txt').write_text(xml_text(ROOT/'A_B_Testing_Hypothesis_Testing_Assignment.docx'),encoding='utf-8')
doc=Document(ROOT/'report/AB_Testing_Assignment_Report.docx')
doc_text=xml_text(ROOT/'report/AB_Testing_Assignment_Report.docx')
pdf=pymupdf.open(ROOT/'report/AB_Testing_Assignment_Report.pdf')
pdf_text='\n'.join(page.get_text() for page in pdf)
(OUT/'docx_actual.txt').write_text(doc_text,encoding='utf-8')
(OUT/'pdf_actual.txt').write_text(pdf_text,encoding='utf-8')
# Exact substantive character stream, ignoring renderer whitespace and repeating furniture.
clean_pdf=re.sub(r'A/B TESTING\s*\|\s*SESSION DURATION\s*Page \d+','',pdf_text)
compact=lambda s: re.sub(r'\s+','',s)
assert compact(doc_text)==compact(clean_pdf), 'DOCX/PDF substantive text differs'
assert len(doc.inline_shapes)==4
assert sum(len(page.get_image_info()) for page in pdf)==4
for j,g in enumerate(['UI_A','UI_B'],1):
    for i,key in enumerate(['count','mean','std','min','25%','50%','75%','max'],1):
        value=desc.loc[g,key]
        assert doc.tables[1].cell(i,j).text==(f'{value:,.0f}' if key=='count' else f'{value:.4f}')
for i,(g,row) in enumerate(normal.iterrows(),1):
    assert [c.text for c in doc.tables[2].rows[i].cells]==[g,str(int(row['count'])),f'{row.sample_skewness:.6f}',f'{row.shapiro_W:.6f}',f'{row.shapiro_p_value:.6e}']
for i,key in enumerate(['mean_difference_seconds','standard_error_seconds','t_statistic','degrees_of_freedom','p_value'],1):
    assert doc.tables[3].cell(i,1).text==format(test[key],'.8e' if key=='p_value' else '.8f')
assert doc.tables[3].cell(6,1).text==f'[{test.ci_95_low_seconds:.8f}, {test.ci_95_high_seconds:.8f}]'
for i,row in probs.iterrows():
    assert [c.text for c in doc.tables[4].rows[i+1].cells][1:]==[f'{row.empirical_count}/{row.denominator}',f'{row.empirical_proportion:.6f}',f'{row.empirical_percentage:.4f}',f'{row.normal_percentage:.4f}']
assert [[c.text for c in r.cells] for r in doc.tables[-1].rows[1:]]==[list(r.values()) for r in rows[:10]]
nb=nbformat.read(ROOT/'AB_Testing_Assignment.ipynb',4)
code=[c for c in nb.cells if c.cell_type=='code']
assert [c.execution_count for c in code]==list(range(1,len(code)+1))
assert not [o for c in code for o in c.outputs if o.output_type=='error']
assert sum('image/png' in o.get('data',{}) for c in code for o in c.outputs)==4
nb_text='\n'.join(c.source if c.cell_type=='markdown' else '\n'.join(o.get('text',o.get('data',{}).get('text/plain','')) for o in c.outputs) for c in nb.cells)
(OUT/'notebook_actual.txt').write_text(nb_text,encoding='utf-8')
# Check every decimal/scientific value in reports/notebook prose and outputs against
# independently calculated values, source data, or explicitly documented constants.
expected_values += [float(v) for r in rows[:10] for k,v in r.items() if k not in ['ui_version','timestamp']]
expected_values += [68.45,71.23,sA-68.45,sB-71.23,.05,.95,.975,.25,.5,.75,.025,1e-10,1e-12,1e-13,1e-8,.00005,26.60,99.66,126.26,0,1,1.37]
unmatched={}
for name,content in [('docx',doc_text),('pdf',pdf_text),('notebook',nb_text)]:
    content=re.sub(r'(?m)^\d+\.\d+ (?=[A-Z])','',content)  # Section numbers, not observations.
    missing=[]; total=0
    for match in re.finditer(r'(?<![\w.])[-+]?\d+\.\d+(?:[eE][-+]?\d+)?(?![\w.])',content):
        token=match.group(); value=float(token)
        mantissa=token.lower().split('e')[0]; exp=int(token.lower().split('e')[1]) if 'e' in token.lower() else 0
        tolerance=.50001*10**(exp-len(mantissa.split('.')[1]))
        if not any(abs(value-v)<=tolerance+abs(v)*1e-13 for v in expected_values):
            missing.append({'token':token,'context':content[max(0,match.start()-45):match.end()+45]})
        total+=1
    unmatched[name]=missing
    checks.append(f'{name}: {total} decimal/scientific occurrences inspected')
(OUT/'numeric_unmatched.json').write_text(json.dumps(unmatched,indent=2),encoding='utf-8')
assert not any(unmatched.values()), unmatched
before=json.loads((OUT/'original_hashes_before.json').read_text(encoding='utf-8-sig'))
for entry in before:
    assert hashlib.sha256(Path(entry['Path']).read_bytes()).hexdigest().upper()==entry['Hash']
summary={'checks':checks,'code_cells':len(code),'pages':len(pdf),'docx_pdf_text_identical_ignoring_whitespace':True,'figures':4,'appendix_cells_verified':90,'probabilities':probability_results,'welch':{'difference':delta,'SE':se,'t':t,'df':df,'p':p,'CI':[lo,hi]},'relative_increase_percent':increase,'original_hashes':before}
(OUT/'verification.json').write_text(json.dumps(summary,indent=2),encoding='utf-8')
print(json.dumps(summary,indent=2))
