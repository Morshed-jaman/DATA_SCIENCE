"""Build the assignment DOCX, export with LibreOffice, and render every PDF page.

Use the project .venv. Numeric report values come from verified results/*.csv.
Preset: standard_business_brief; compact memo masthead without a cover page.
Named overrides: 10pt table/caption text; 9pt landscape data table; 22pt title.
"""
from pathlib import Path
import argparse
import hashlib
import json
import os
import shutil
import subprocess
import sys

import numpy as np
import pandas as pd
import nbformat
import pymupdf
from scipy import stats
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'report'
NAME = 'AB_Testing_Assignment_Report'


def load(name, **kwargs):
    return pd.read_csv(ROOT / 'results' / name, float_precision='round_trip', **kwargs)


def build():
    OUT.mkdir(exist_ok=True)
    raw = pd.read_csv(ROOT / 'ui_ab_testing_data.csv')
    desc = load('descriptive_statistics.csv', index_col='ui_version')
    normal = load('normality_assessment.csv', index_col='ui_version')
    variance = load('variance_test.csv').iloc[0]
    test = load('t_test_results.csv').query("role == 'selected'").iloc[0]
    probs = load('probability_scenarios.csv')
    business = load('business_impact.csv').iloc[0]
    # Read the original assignment, including table cells, not only its extraction.
    original = Document(ROOT / 'A_B_Testing_Hypothesis_Testing_Assignment.docx')
    assignment_text = '\n'.join([p.text for p in original.paragraphs] +
        [cell.text for table in original.tables for row in table.rows for cell in row.cells])
    assert 'Probability' in assignment_text and len(original.tables) == 2
    notebook = nbformat.read(ROOT / 'AB_Testing_Assignment.ipynb', as_version=4)
    for cell in notebook.cells:
        if cell.cell_type == 'code':
            assert cell.execution_count is not None
            assert not any(o.output_type == 'error' for o in cell.outputs)
    (ROOT / 'requirements_checklist.md').read_text(encoding='utf-8')
    for label in desc.index:
        values = raw.loc[raw.ui_version.eq(label), 'session_duration_seconds']
        np.testing.assert_allclose([len(values), values.mean(), values.std(ddof=1)],
                                  desc.loc[label, ['count', 'mean', 'std']], rtol=1e-13, atol=0)
    A, B = desc.loc['UI_A'], desc.loc['UI_B']
    assert test.test_name == 'Welch' and variance.p_value < .05
    np.testing.assert_allclose(test.mean_difference_seconds, B['mean']-A['mean'], rtol=1e-13)
    np.testing.assert_allclose(business.percentage_increase,
                              100*(B['mean']-A['mean'])/A['mean'], rtol=1e-13)
    np.testing.assert_allclose(test.p_value, 2*stats.t.sf(abs(test.t_statistic), test.degrees_of_freedom), rtol=1e-10, atol=0)
    for _, r in probs.iterrows():
        d = stats.norm(r.mean_seconds, r.sample_sd_seconds)
        expected = d.sf(r.lower_bound_seconds) if np.isposinf(r.upper_bound_seconds) else d.cdf(r.upper_bound_seconds)-d.cdf(r.lower_bound_seconds)
        np.testing.assert_allclose(expected, r.normal_probability, rtol=1e-13, atol=0)
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = False
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(.492)
    for name, size, before, after, color in [
        ('Normal',11,0,6,'20252B'),('Title',22,0,5,'17324D'),
        ('Subtitle',12,0,10,'555555'),('Heading 1',16,16,8,'2E74B5'),
        ('Heading 2',13,12,6,'2E74B5'),('Heading 3',12,8,4,'1F4D78'),
        ('Caption',10,4,4,'555555')]:
        style = doc.styles[name]
        style.font.name = 'Calibri'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1
        if name == 'Caption': style.font.bold = False
        border = style.element.find('.//' + qn('w:pBdr'))
        if border is not None: border.getparent().remove(border)
    header = section.header.paragraphs[0]
    header.text = 'A/B TESTING  |  SESSION DURATION'
    header.runs[0].font.size = Pt(9)
    header.runs[0].font.color.rgb = RGBColor.from_string('666666')
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Page ').font.size = Pt(9)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    def p(text, style=None):
        return doc.add_paragraph(text, style)

    def h(text):
        doc.add_heading(text, level=1)

    def page(text):
        heading = doc.add_heading(text, level=1)
        heading.paragraph_format.page_break_before = True

    def eq(text):
        para = doc.add_paragraph()
        para.paragraph_format.space_after = Pt(7)
        math = OxmlElement('m:oMath')
        run = OxmlElement('m:r'); txt = OxmlElement('m:t'); txt.text = text
        run.append(txt); math.append(run); para._p.append(math)
        return para

    def table(headers, rows, widths=None, size=10):
        widths = widths or [6.5/len(headers)]*len(headers)
        twips = [round(w*1440) for w in widths]
        t = doc.add_table(rows=1, cols=len(headers)); t.style = 'Table Grid'; t.autofit = False
        pr = t._tbl.tblPr
        pr.find(qn('w:tblW')).set(qn('w:w'), str(sum(twips)))
        pr.find(qn('w:tblW')).set(qn('w:type'), 'dxa')
        indent = OxmlElement('w:tblInd'); indent.set(qn('w:w'),'120'); indent.set(qn('w:type'),'dxa'); pr.append(indent)
        margins = OxmlElement('w:tblCellMar')
        for side,value in [('top',80),('bottom',80),('start',120),('end',120)]:
            node=OxmlElement('w:'+side); node.set(qn('w:w'),str(value)); node.set(qn('w:type'),'dxa'); margins.append(node)
        pr.append(margins)
        for col,w in zip(t.columns,twips): col.width=Inches(w/1440)
        for cell,text in zip(t.rows[0].cells,headers): cell.text=str(text)
        repeat=OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
        for row in rows:
            for cell,text in zip(t.add_row().cells,row): cell.text=str(text)
        for i,row in enumerate(t.rows):
            no_split=OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
            for j,cell in enumerate(row.cells):
                cell.width=Inches(twips[j]/1440)
                cell.vertical_alignment=1
                if i==0:
                    shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'F2F4F7'); cell._tc.get_or_add_tcPr().append(shade)
                for para in cell.paragraphs:
                    para.paragraph_format.space_after=Pt(2)
                    para.paragraph_format.line_spacing=1.0
                    for run in para.runs: run.font.size=Pt(size); run.bold=(i==0)
        return t

    def figure(file, caption):
        para=doc.add_paragraph(); para.paragraph_format.keep_with_next=True
        para.alignment=WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(str(ROOT/'figures'/file),width=Inches(5.65))
        p(caption,'Caption')

    p('A/B Testing and Hypothesis Testing','Title')
    p('Comparing Session Duration for UI_A and UI_B','Subtitle')
    p('Name: ____________________    Student ID: ____________________')
    p('Course: ___________________    Section: _______________________')
    h('1. Objective and hypotheses')
    p('Compare population mean session duration between the two mobile-app interfaces. '
      'The target metric is session_duration_seconds, measured in seconds. All differences '
      'are defined as UI_B minus UI_A. Retain every observation without transformations.')
    eq('H₀: μB − μA = 0     H₁: μB − μA ≠ 0     α = 0.05 (two-sided)')
    h('2. Data exploration')
    p(f'The supplied CSV contains {len(raw):,} session records and {len(raw.columns)} columns '
      f'({len(raw)+1:,} records including the header). The first ten rows appear in Appendix A.')
    descriptions = [
        ('user_id','Session identifier','Integer'),('ui_version','Interface shown: UI_A or UI_B','String'),
        ('timestamp','Session date and time; loaded as text','Date/time text'),
        ('session_duration_\nseconds','Session length; target metric, seconds','Float'),
        ('pages_visited','Pages/screens viewed','Integer'),('clicks','Total interactions','Integer'),
        ('errors_encountered','Errors or bugs experienced','Integer'),
        ('conversion','Desired action completed: 1=yes, 0=no','Binary'),
        ('bounce','Left without interaction: 1=yes, 0=no','Binary')]
    table(['CSV column','Description','Type'],descriptions,[2.15,3.25,1.1])
    p('Column names above preserve CSV capitalization. The assignment capitalizes four '
      'headers differently; the analysis uses the actual CSV names.','Caption')

    page('2.1 Data quality and descriptive statistics')
    from analyze_ab_testing import quality_checks
    checks=quality_checks(raw)
    p('Data-quality checks: ' + '; '.join(f'{k.replace("_", " ")} = {int(v)}' for k,v in checks.items()) + '.')
    p('Every column has zero missing values. No records were removed, imputed, or transformed. '
      'Distinct session IDs do not demonstrate independent users or observations.')
    table(['Statistic','UI_A','UI_B'],[(label, f'{A[key]:,.0f}' if key=='count' else f'{A[key]:.4f}',
        f'{B[key]:,.0f}' if key=='count' else f'{B[key]:.4f}') for key,label in
        [('count','Sample size'),('mean','Mean (seconds)'),('std','Sample SD (seconds)'),
         ('min','Minimum (seconds)'),('25%','25th percentile (seconds)'),('50%','Median (seconds)'),
         ('75%','75th percentile (seconds)'),('max','Maximum (seconds)')]],[3.3,1.6,1.6])
    p('Sample SD uses ddof=1; quartiles use linear interpolation. SD describes variation '
      'among individual durations. Standard error describes uncertainty in an estimated '
      'mean (SD/√n under independence). These quantities are not interchangeable.')
    p(f'The assignment preview lists SDs of 68.45 and 71.23 seconds, whereas the CSV yields '
      f'{A["std"]:.4f} and {B["std"]:.4f}. Counts and rounded means match. All calculations use the CSV.')
    h('3. Normality assessment')
    table(['Group','n','Skewness¹','Shapiro W','p-value'],[[g,int(r['count']),f'{r.sample_skewness:.6f}',
        f'{r.shapiro_W:.6f}',f'{r.shapiro_p_value:.6e}'] for g,r in normal.iterrows()], [1,.7,1.3,1.3,2.2])
    p('¹ Sample skewness uses bias=False. Shapiro H₀: normal distribution; H₁: non-normal. '
      'Both p-values are below 0.05, supporting rejection. A large p-value would not prove normality.','Caption')

    page('3.1 Histograms of individual durations')
    figure('UI_A_histogram.png','Figure 1. UI_A: 25 bins; a clear right tail and asymmetric concentration at shorter durations.')
    figure('UI_B_histogram.png','Figure 2. UI_B: the same 25 bin boundaries and axis limits. The distribution is right-skewed and extends farther to the right.')
    p('Both are unimodal but not symmetric bell shapes. All observations, including extreme durations, remain visible in the full-range plots.')

    page('3.2 Normal Q-Q plots and inference')
    figure('UI_A_qq_plot.png','Figure 3. UI_A normal probability plot (scipy.stats.probplot), with its fitted reference line.')
    figure('UI_B_qq_plot.png','Figure 4. UI_B normal probability plot, with its fitted reference line. Both show systematic curvature and large upper-tail departures.')
    p('The lower tails lie above the fitted lines, central portions tend below them, and '
      'upper tails rise sharply above them. Positive skewness corroborates these departures.')

    page('3.3 Normality limitations and assumptions')
    p('UI_A uses all 5,013 observations. SciPy issued the following warning; no subsampling was performed:')
    p(str(normal.loc['UI_A','warning']))
    p('For more than 5,000 observations, SciPy documents that W remains accurate but the '
      'p-value may not. Treat its precise magnitude cautiously. UI_B has 4,987 observations '
      'and did not trigger this warning. Large samples make normality tests sensitive, but '
      'the observed curvature and skewness also show substantial non-normality [1].')
    p('The distribution of individual durations is not the sampling distribution of a '
      'mean. Under independence and finite population variances, large samples support '
      'approximate mean inference even with non-normal raw data. Independence is assumed '
      'from the assignment scenario; unique IDs do not establish it. Large samples do not '
      'repair dependence or biased sampling. Welch addresses unequal variances, not raw-data normality.')
    h('4. Variance assessment and Welch test')
    p(f'Median-centered Levene (Brown-Forsythe) gives statistic {variance.statistic:.6f}, '
      f'p = {variance.p_value:.8e}. It is robust to departures from normality, especially '
      'skewness [2]. The assignment rule selects Welch because p < 0.05. A non-significant '
      'result would not prove equal variances; preliminary variance-based selection is '
      'data-dependent and can affect overall error properties.')
    p('Let mA, mB be sample means; sA, sB sample SDs; and nA, nB sample sizes. Define vA = sA²/nA and vB = sB²/nB.')
    eq('SE = √(vB + vA);   t = (mB − mA)/SE')
    eq('df = (vB + vA)² / [vB²/(nB − 1) + vA²/(nA − 1)]')
    eq('p = 2 × t.sf(|t|, df)')
    eq('CI₉₅ = (mB − mA) ± t.ppf(0.975, df) × SE')
    vB, vA=B['std']**2/B['count'],A['std']**2/A['count']
    p(f'vB = {B["std"]:.8f}²/{int(B["count"])} = {vB:.10f}; '
      f'vA = {A["std"]:.8f}²/{int(A["count"])} = {vA:.10f}.')
    p(f'SE = √({vB:.10f} + {vA:.10f}) = {test.standard_error_seconds:.8f} seconds. '
      f't = ({B["mean"]:.8f} − {A["mean"]:.8f})/{test.standard_error_seconds:.8f} = {test.t_statistic:.8f}.')
    p(f'df = ({vB:.8f} + {vA:.8f})² / [{vB:.8f}²/{int(B["count"])-1} + '
      f'{vA:.8f}²/{int(A["count"])-1}] = {test.degrees_of_freedom:.8f}.')

    page('4.1 Test results and interpretation')
    table(['Welch result: UI_B minus UI_A','Value'],[
        ('Mean difference (seconds)',f'{test.mean_difference_seconds:.8f}'),
        ('Standard error (seconds)',f'{test.standard_error_seconds:.8f}'),
        ('t-statistic',f'{test.t_statistic:.8f}'),('Degrees of freedom',f'{test.degrees_of_freedom:.8f}'),
        ('Two-sided p-value',f'{test.p_value:.8e}'),
        ('95% CI (seconds)',f'[{test.ci_95_low_seconds:.8f}, {test.ci_95_high_seconds:.8f}]'),
        ('Decision at α = 0.05',test.decision)],[3.4,3.1])
    critical=stats.t.ppf(.975,test.degrees_of_freedom)
    p(f'CI = {test.mean_difference_seconds:.8f} ± {critical:.8f} × '
      f'{test.standard_error_seconds:.8f} = [{test.ci_95_low_seconds:.8f}, '
      f'{test.ci_95_high_seconds:.8f}] seconds.')
    p(f'Reject H₀. UI_B has a higher estimated mean duration by {test.mean_difference_seconds:.4f} '
      'seconds. The confidence interval excludes zero. The p-value describes tail extremeness '
      'under H₀ and the model; it is not the probability that H₀ is true. Statistical significance '
      'does not by itself establish business value.')
    p('The raw-data SciPy test uses UI_B first, UI_A second, equal_var=False, and a two-sided '
      'alternative [3]. Independent summary-statistic formulas agree with SciPy, including '
      'its confidence interval. Tiny p-values were checked with relative tolerance 1e-10, '
      'absolute tolerance zero, and positivity assertions; the script and notebook also agree.')
    h('5. Required normal-model probability scenarios')
    eq('XA ~ Normal(mA, sA²);   XB ~ Normal(mB, sB²)')
    p(f'mA = {A["mean"]:.8f}, sA = {A["std"]:.8f}; mB = {B["mean"]:.8f}, '
      f'sB = {B["std"]:.8f} seconds. Use individual-session SDs, not standard errors. '
      'All displayed substitutions are rounded only for readability.')
    p('These are the normal-model approximations required by the assignment. Both raw '
      'distributions are right-skewed and non-normal; large samples do not make individual '
      'durations normal. The requested model is not truncated, even though it allows '
      'negative theoretical durations. Empirical comparisons are reported separately.')

    page('5.1 Worked probability answers')
    events=['P(XA > mB)','P(XA < 80)','P(XB > 150)','P(80 < XA < 120)']
    for i,r in probs.iterrows():
        doc.add_heading(f'({r.question}) {events[i]}',level=2)
        p(r.event+'.')
        for side in ['lower','upper']:
            bound=r[f'{side}_bound_seconds']
            if np.isfinite(bound):
                p(f'z{side} = ({bound:.8f} − {r.mean_seconds:.8f}) / {r.sample_sd_seconds:.8f} = {r[f"z_{side}"]:.8f}.')
        if np.isposinf(r.upper_bound_seconds):
            formula=f'1 − Φ(z) = norm.sf({r.z_lower:.8f})'
        elif np.isneginf(r.lower_bound_seconds):
            formula=f'Φ(z) = norm.cdf({r.z_upper:.8f})'
        else:
            formula=f'Φ(zupper) − Φ(zlower) = Φ({r.z_upper:.8f}) − Φ({r.z_lower:.8f})'
        p(f'{formula} = {r.normal_probability:.8f} = {r.normal_percentage:.4f}%.')
        p(f'Under this model, about {r.normal_percentage:.2f}% of individual {r.group} sessions satisfy the event.')
        if i==0: p('UI_B’s full-precision sample mean is a fixed threshold here. This is not P(XA > XB).')

    page('5.2 Supplementary empirical comparison')
    p('Observed proportions use the matching group denominator, strict inequalities, '
      'and the full-precision UI_B mean for (a). These sample proportions supplement, '
      'rather than replace, the required normal-model answers.')
    table(['Event','Count / n','Observed proportion','Observed %','Model %'],[
        [events[i],f'{r.empirical_count}/{r.denominator}',f'{r.empirical_proportion:.6f}',
         f'{r.empirical_percentage:.4f}',f'{r.normal_percentage:.4f}'] for i,r in probs.iterrows()],
        [1.35,1.15,1.5,1.25,1.25])
    p('The normal approximation differs from the empirical distribution: right skew '
      'changes the allocation of mass around the thresholds and in the tails. Empirical '
      'rates describe this sample, not exact population probabilities. The strongest '
      'mismatches occur in (a), (b), and (c).')
    h('6. Business impact and conclusion')
    eq('Absolute difference = mB − mA')
    p(f'= {business.mean_B_seconds:.8f} − {business.mean_A_seconds:.8f} = '
      f'{business.absolute_difference_seconds:.8f} seconds.')
    eq('Relative increase (%) = 100 × (mB − mA)/mA')
    p(f'= 100 × ({business.mean_B_seconds:.8f} − {business.mean_A_seconds:.8f}) / '
      f'{business.mean_A_seconds:.8f} = {business.percentage_increase:.6f}%.')
    p('This is a relative percentage increase in mean session duration, not a change '
      'in percentage points. The verified Welch result is statistically significant '
      f'(p = {test.p_value:.8e}); its 95% interval for UI_B minus UI_A is '
      f'[{test.ci_95_low_seconds:.6f}, {test.ci_95_high_seconds:.6f}] seconds.')
    p('If longer engagement is the business objective, UI_B is a candidate for adoption, '
      'conditional on practical value, costs, and checks of relevant outcomes. Longer '
      'sessions alone do not prove better usability, satisfaction, conversion, or revenue; '
      'they may also reflect friction. The evidence establishes a difference in mean '
      'duration under the stated assumptions, not a universal improvement in user experience.')

    page('7. Reproducibility and sources')
    p('Supplied sources: A_B_Testing_Hypothesis_Testing_Assignment.docx and '
      'ui_ab_testing_data.csv. Original files remain unchanged. Numeric tables and '
      'substitutions in this report are loaded from the verified analysis outputs; '
      'figures are the previously inspected exports.')
    p('From the project folder, follow README.md to use the Python 3.13.1 .venv and '
      'requirements.txt. Run the analysis script, then the full notebook in a fresh '
      'project kernel. Finally run the report builder. The accompanying notebook and '
      'script retain full code; this report presents methods and findings.')
    for command in ['.\\.venv\\Scripts\\python.exe scripts\\analyze_ab_testing.py',
                    '.\\.venv\\Scripts\\python.exe scripts\\execute_notebook.py',
                    '.\\.venv\\Scripts\\python.exe scripts\\build_report.py']:
        para=p(command); para.runs[0].font.name='Consolas'; para.runs[0].font.size=Pt(9)
    p('Saved analysis sources: descriptive_statistics.csv, normality_assessment.csv, '
      'variance_test.csv, t_test_results.csv, probability_scenarios.csv, and '
      'business_impact.csv in results/. Full precision is retained in those files; '
      'only reported displays are rounded. Tiny p-values use scientific notation.')
    p('Checks include normal-probability z-score versus direct loc/scale calculations, '
      'strict empirical counts and group denominators, probability bounds, and '
      'script/notebook/manual-formula agreement. Report generation checks key values '
      'against the source CSV and saved results without replacing them.')
    doc.add_heading('Documentation consulted',level=2)
    sources=[('1','Shapiro-Wilk: sample-size warning','shapiro'),
             ('2','Levene: median-centered variance assessment','levene'),
             ('3','Independent t-test: Welch and confidence interval','ttest_ind'),
             ('4','Probability plots: normal quantiles and fitted line','probplot')]
    for num,title,slug in sources:
        p(f'[{num}] SciPy documentation. {title}.')
        para=p(f'https://docs.scipy.org/doc/scipy/reference/generated/scipy.stats.{slug}.html')
        para.runs[0].font.size=Pt(9)
    p('The normal Q-Q visualizations use scipy.stats.probplot with normal quantiles '
      'and a fitted reference line [4]. No external datasets or invented references are used.')

    landscape=doc.add_section(WD_SECTION_START.NEW_PAGE)
    landscape.orientation=WD_ORIENT.LANDSCAPE
    landscape.page_width,landscape.page_height=Inches(11),Inches(8.5)
    h('Appendix A. First ten CSV records')
    p('All nine columns are shown in their original row order. Short display headers '
      'are mapped below to the exact CSV names; all durations are in seconds.')
    headers=['user_id','UI','timestamp','Duration\n(seconds)','Pages','Clicks','Errors','Conversion','Bounce']
    table(headers,raw.head(10).astype(str).values.tolist(),[.6,.6,1.5,1.2,.75,.75,.75,1, .85],size=9)
    p('UI = ui_version; Duration = session_duration_seconds; Pages = pages_visited; '
      'Clicks = clicks; Errors = errors_encountered; Conversion = conversion; '
      'Bounce = bounce. user_id and timestamp retain their original names.','Caption')
    p('The timestamp is displayed exactly as loaded. Conversion and bounce are binary '
      'indicators defined in Section 2. Their business outcomes were not analyzed in this assignment.')
    doc.core_properties.title='A/B Testing and Hypothesis Testing'
    doc.core_properties.subject='Comparing Session Duration for UI_A and UI_B'
    doc.core_properties.author=''
    target=OUT/f'{NAME}.docx'; doc.save(target)
    return target


def render(target, renderer=None):
    candidates=[renderer,os.environ.get('SOFFICE'),shutil.which('soffice'),
                str(ROOT/'.tools'/'LibreOffice'/'program'/'soffice.exe')]
    executable=next((x for x in candidates if x and Path(x).is_file()),None)
    if not executable:
        raise RuntimeError('DOCX saved; LibreOffice renderer unavailable. Set SOFFICE or pass --soffice.')
    profile=OUT/'qa'/'lo_profile'; profile.mkdir(parents=True,exist_ok=True)
    subprocess.run([executable,f'-env:UserInstallation={profile.as_uri()}', '--headless',
                    '--convert-to','pdf','--outdir',str(OUT),str(target)],check=True,timeout=180)
    pdf=OUT/f'{NAME}.pdf'
    if not pdf.is_file(): raise RuntimeError('Renderer did not create the PDF.')
    document=pymupdf.open(pdf)
    qa=OUT/'qa'; qa.mkdir(exist_ok=True)
    page_text=[]
    for i,page in enumerate(document):
        page.get_pixmap(matrix=pymupdf.Matrix(1.3,1.3)).save(qa/f'page-{i+1:02d}.png')
        page_text.append(page.get_text())
    (qa/'rendered_text.txt').write_text('\n\n'.join(page_text),encoding='utf-8')
    (qa/'render_metadata.json').write_text(json.dumps({'docx_rendered_page_count':len(document),
        'pdf_page_count':len(document),'renderer':executable,'visual_inspection':'pending'},indent=2),encoding='utf-8')
    print(f'Generated DOCX and PDF; {len(document)} pages rendered for inspection.')


if __name__=='__main__':
    parser=argparse.ArgumentParser(); parser.add_argument('--soffice'); parser.add_argument('--docx-only',action='store_true')
    args=parser.parse_args()
    path=build()
    if not args.docx_only: render(path,args.soffice)
    print(path)
